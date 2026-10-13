from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "authority" / "output" / "念念智剪权威项目文档V2.docx"
VIS = ROOT / "docs" / "authority" / "visuals"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def set_cell_text(cell, text, bold=False, color="1F2937", size=8.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)

def add_body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.12
    p.add_run(text)

def add_stage(doc, title, body):
    doc.add_heading(title, level=2)
    parts = body.split("。")
    for part in parts:
        part = part.strip()
        if not part:
            continue
        if "：" in part and len(part) < 220:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
            key, value = part.split("：", 1)
            r = p.add_run(key + "：")
            r.bold = True
            p.add_run(value)
        else:
            add_body(doc, part + "。")

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.65)
sec.bottom_margin = Inches(0.65)
sec.left_margin = Inches(0.72)
sec.right_margin = Inches(0.72)

styles = doc.styles
styles["Normal"].font.name = "Microsoft YaHei"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
styles["Normal"].font.size = Pt(10)
for name, size, color in [("Title", 28, "111827"), ("Heading 1", 18, "111827"), ("Heading 2", 13, "C2413B"), ("Heading 3", 11, "374151")]:
    styles[name].font.name = "Microsoft YaHei"
    styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles[name].font.size = Pt(size)
    styles[name].font.color.rgb = RGBColor.from_string(color)

header = sec.header.paragraphs[0]
header.text = "念念智剪  |  权威项目文档 V2  |  前期准备，不发布"
header.style = styles["Normal"]
header.runs[0].font.size = Pt(8)
header.runs[0].font.color.rgb = RGBColor(107, 114, 128)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("内部产品与工程基线 · 2026-08-10").font.size = Pt(8)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("念念智剪权威项目文档 V2")
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.add_run("OpenChatCut 底座上的念念品牌剪辑器 · 前期产品与工程主控基线").italic = True
add_body(doc, "本文件只描述念念智剪本体，不吞并念念 AI 主站或念念智选业务。它把“先看到最终效果，再开始构建”固定为每个阶段的工作顺序；本版本只做权威准备，不发布线上。")

doc.add_heading("一页看懂", level=1)
add_body(doc, "项目定位：念念智剪是两个宿主共用的真实视频剪辑能力。念念 AI 负责 AI 转绘和 AI 短剧，念念智选负责电商带货；两个宿主都可以启动智剪成片方案，智剪只负责工程、媒体、剪辑、质检和交付。")
add_body(doc, "最高价值最小闭环：用户从快速模式选择“无声素材解说”，上传一段视频并填写主题，系统生成真实口播、真实词级字幕、BGM 和可编辑工程，再一键进入专业模式继续修改并导出。")
add_body(doc, "本文件的通过标准：负责人能先看懂 5 张目标图；每个 S0–S7 有唯一交付物和证据；Codex 能按提示词逐阶段执行；任何未验证的 Provider、线上变更或用户媒体都不被文档伪造为完成。")

doc.add_heading("S0 基线与边界", level=1)
baseline = [
    "项目目录：E:\\codex\\niannianai\\niannianzhijian；线上：https://edit.cau.fun/。",
    "Git origin：1008611-creater/niannianzhijian；upstream：0xsline/OpenChatCut；当前分支 codex/qwen-forced-aligner-cuda126。",
    "当前线上发行：r21-hosted-admin-billing-20260810；r20 保留回滚；本次只读确认，不发布。",
    "本地入口：bash local-wsl-dev/start-local.sh；验证：bash local-wsl-dev/verify-local.sh；停止：bash local-wsl-dev/stop-local.sh；端口 3026/5199。",
    "必须保留：MiMo ASR/TTS、自定义供应商、Qwen3 ForcedAligner、SSO/积分边界、预览代理、上传/导出修复、念念 Logo、AGPL 开源入口。",
    "缺失文件：项目根目录未找到 NIANNIAN_CUSTOMIZATION_BOUNDARY.md；本 V2 以 AGENTS.md、DESIGN.md、迭代计划和 FAST_PRO_MODE_PLAN_ZH.md 的边界为权威，并把该缺失记录为文档缺口。",
]
for item in baseline:
    add_bullet(doc, item)

doc.add_heading("编辑器对标结论", level=1)
add_body(doc, "CapCut：借鉴“先选成片目标再进入编辑器”的入口；念念差异是方案卡要说明素材、结果和下一步，并且结果保留真实工程。")
add_body(doc, "Descript：借鉴转写/文本驱动修改；念念差异是把真实词级时间戳、字幕和音频关系写进工程，MiMo 整段文本不能冒充词级字幕。")
add_body(doc, "OpusClip：借鉴长视频到发布版的结果预期、平台规格和可分享交付；念念差异是输出仍可回到专业模式，删除建议需要用户确认。")

doc.add_heading("目标界面图", level=1)
add_body(doc, "以下图只规定视觉方向、信息层级和用户下一步，不代表功能已实现。图片以当前 edit.cau.fun 页面和智剪深色时间线风格为基线制作；Image2 上游本轮返回 502，服务恢复后可以按同一结构替换，不改变合同。")
figs = [
    ("01-quick-home.png", "快速模式首页：用户知道快速成片与专业剪辑的区别，并能按结果选择。"),
    ("02-recipe-select.png", "成片方案选择：明确需要什么素材、会得到什么、下一步上传什么。"),
    ("03-quick-running.png", "快速运行状态：显示阶段进度、可恢复边界和完成后进入专业模式。"),
    ("04-professional-editor.png", "专业编辑器：素材、预览、字幕、配音、BGM、Agent 和时间线在同一真实工程。"),
    ("05-export-delivery.png", "导出/交付页：质检结果、文件规格、下载和回到专业编辑的动作清楚。"),
]
for fname, caption in figs:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(VIS / fname), width=Inches(6.65))
    cp = doc.add_paragraph(caption)
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.runs[0].italic = True
    cp.runs[0].font.size = Pt(9)
    doc.add_page_break()

doc.add_heading("双模式信息架构", level=1)
add_body(doc, "快速模式首屏只问“你想做哪种视频？”，用户看到“成片方案/视频类型/发布场景”，不看到 API、模型、供应商、Agent 工具名或时间线内部术语。每张方案卡必须写成片结果、适用场景、素材要求、预计时长/平台、字幕/配音情况和下一步按钮。")
add_body(doc, "专业模式保留当前完整编辑器。快速模式写入的片段、字幕、旁白、BGM、包装对象必须能继续编辑；两种模式共用同一个 projectId 和真实工程，不复制第二套时间线逻辑。")
add_body(doc, "宿主默认排序：念念 AI 为短剧片段精修、剧情混剪、转绘素材成片；念念智选为商品种草、测评口播、卖点讲解；两者共用无声素材解说。")

doc.add_heading("S0–S7 阶段总表", level=1)
rows = [
    ("S0", "运行基线", "确认 upstream/origin、r21/r20、WSL 和念念边界", "chatcut-plugin-basics", "remote/status/log、候选目录、线上只读回读"),
    ("S1", "双模式蓝图", "冻结入口、状态和同一真实工程关系", "frontend-taste-router", "页面合同、路径回放、目标图审阅"),
    ("S2", "方案注册表", "按宿主定义可理解的成片方案", "frontend-taste-router", "字段测试、宿主排序、UI 不泄漏内部词"),
    ("S3", "真实运行契约", "素材到导出的可恢复编排", "按缺口单选媒体 Skill", "工程 JSON、真实时间戳、阶段重试"),
    ("S4", "专业模式", "快速对象进入完整时间线", "chatcut-plugin-basics", "同一 projectId、对象可见可改"),
    ("S5", "编辑闭环", "导入、对齐、清洗建议、字幕、声音和真实导出", "按缺口单选媒体 Skill", "音画同步、无黑帧、字幕覆盖"),
    ("S6", "宿主接入", "AI/智选通过 project/asset/delivery 合同接入", "ai-native-collaboration", "两个真实入口与返回编辑"),
    ("S7", "发布验证", "形成可发布候选，本轮不发布", "browser-control-in-app-browser", "WSL、浏览器、版本哈希、AGPL 清单"),
]
table = doc.add_table(rows=1, cols=5)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Table Grid"
for i, h in enumerate(["阶段", "主题", "目标", "主 Skill", "验收证据"]):
    set_cell_text(table.rows[0].cells[i], h, bold=True, color="FFFFFF", size=8)
    shade(table.rows[0].cells[i], "C2413B")
for row in rows:
    cells = table.add_row().cells
    for i, value in enumerate(row):
        set_cell_text(cells[i], value, size=7.5)

stage_text = {
"S0": "目标：确认 OpenChatCut 与公网运行基线。真实输入：Git remote/status/log、现有项目文档、local-wsl-dev、非敏感线上状态。唯一交付物：基线与冲突清单。执行方案：只读检查，不拉取、不覆盖、不发布。主 Skill：chatcut:chatcut-plugin-basics 的项目路径与底座章节。支持 Skill：chatcut:known-errors，仅在复现 RefreshReg、上传或空间错误时使用。Agent：主控负责边界；上游同步 Agent 只写对照清单。用户结果：知道哪些能力可用、哪些改动不能被上游覆盖。验收证据：remote、提交、候选目录、线上回读、本地入口。下一阶段输入：冻结基线。",
"S1": "目标：定义快速成片与专业剪辑的入口、状态和同工程关系。真实输入：S0、5 张目标图、编辑器对标观察。唯一交付物：双模式页面合同和用户路径图。执行方案：先冻结界面状态与文案，再确定对象映射；快速结果回到专业模式。主 Skill：frontend-taste-router 的 Foundation / Design Read / Required Work Cycle。支持 Skill：design-taste-frontend 的 Brief Inference / Anti-default discipline。Agent：主控定边界；快速模式 Agent 维护模式合同；专业编辑器 Agent 维护时间线合同。验收证据：桌面/移动截图、路径回放、同一 projectId。下一阶段输入：模式合同。",
"S2": "目标：建立按宿主排序的方案注册表。真实输入：S1、两个宿主上下文。唯一交付物：recipe registry 与方案卡数据。执行方案：AI 用短剧/剧情/转绘方案，智选用商品/测评/卖点方案，共用无声解说；每卡写结果、场景、素材、时长平台和下一步。主 Skill：frontend-taste-router 的 Quality Rules。支持 Skill：insight-to-skill-compounding 仅在明确批准后沉淀。Agent：快速模式 Agent 唯一维护 registry。验收证据：字段测试、宿主排序、不出现 API/模型词。下一阶段输入：已审阅 registry。",
"S3": "目标：跑成真实可恢复工程。真实输入：recipe、projectId、素材源版本、主题。唯一交付物：运行契约与阶段状态机。执行方案：导入 -> 转写/脚本 -> 口播 -> 字幕 -> BGM -> 包装 -> 验证 -> 导出；记录 recipeId、版本、styleId、workflowRunId、源版本、真实时间戳、创建对象和失败阶段；重跑只替换本次对象。主 Skill：按缺口单选 asset-import / transcription / voice / music / create-motion-graphics / export。支持 Skill：verification。Agent：快速模式只编排，媒体链路只维护媒体。验收证据：工程 JSON、源版本、真实时间戳、合成画面、重试记录。下一阶段输入：可打开快速工程。",
"S4": "目标：确认项目列表、素材库、预览、时间线、Agent、字幕、配音、BGM、视频生成和导出同工程可用。真实输入：S3 工程。唯一交付物：专业能力矩阵与对象映射。执行方案：只补真实缺失，不重写时间线。主 Skill：chatcut:chatcut-plugin-basics 的项目路径章节。支持 Skill：browser-control-in-app-browser 的真实回读。Agent：专业编辑器 Agent 唯一维护时间线。验收证据：同一 projectId、对象可见、可撤销、浏览器回放。下一阶段输入：能力矩阵。",
"S5": "目标：完成素材导入、ASR/强制对齐、口播清洗建议、字幕/配音/音乐、动效和真实导出。真实输入：S4 工程和用户确认修改。唯一交付物：可确认编辑计划与可播放导出。执行方案：删除/清洗先展示，确认后写时间线；MiMo 整段文本不能冒充词级字幕。主 Skill：按缺口单选媒体 Skill。支持 Skill：verification。Agent：媒体链路维护导入/对齐/配音/音乐/导出；禁止静默删除。验收证据：词级覆盖、无黑帧冻结帧长静音、字幕安全区、音画同步、文件回读。下一阶段输入：已验证闭环。",
"S6": "目标：两个宿主通过明确合同启动方案并返回智剪。真实输入：registry、S5 工程、project/asset/delivery 合同。唯一交付物：两个宿主适配器和返回路径。执行方案：宿主只给业务上下文和素材引用，智剪不读宿主数据库。主 Skill：ai-native-collaboration 的职责边界与执行闭环。支持 Skill：asset-import 的引用素材章节。Agent：宿主集成 Agent 唯一维护合同。验收证据：两个真实入口、字段回读、返回编辑、权限积分边界。下一阶段输入：宿主验收记录。",
"S7": "目标：形成可发布候选但本轮不发布。真实输入：S6 完整候选和验收证据。唯一交付物：本地候选清单、发布/回滚清单和 AGPL 发行包边界。执行方案：WSL 构建、verify-local、浏览器首屏/预览/导出/控制台、品牌/SSO/积分/版本哈希/回滚；包排除媒体、env、凭据。主 Skill：browser-control-in-app-browser 的编辑器验收。支持 Skill：verification 的导出质检。Agent：独立验证 Agent 仅候选冻结后检查，发布由主控。验收证据：候选哈希、服务健康、页面回读、真实导出、控制台、AGPL 清单。下一阶段输入：负责人发布许可。",
}
for key in ["S0","S1","S2","S3","S4","S5","S6","S7"]:
    add_stage(doc, key + " 阶段执行说明", stage_text[key])

doc.add_heading("Agent 责任与唯一写入者", level=1)
add_body(doc, "主控负责产品边界、阶段推进和发布；上游同步 Agent 只做 upstream 对照与冲突清单；快速模式 Agent 只维护方案契约与运行编排；专业编辑器 Agent 只维护时间线；媒体链路 Agent 维护导入/转写/配音/音乐/导出；宿主集成 Agent 维护 AI/智选契约；独立验证 Agent 只在候选冻结后检查真实浏览器。Codex 主控是本项目唯一工程写入者，其他 Agent 只能提出补丁或报告。")

doc.add_heading("运行与发布验收矩阵", level=1)
matrix = [
    ("本地运行", "WSL 启停与 verify-local", "服务活跃、3026/5199 可读、项目重启不丢"),
    ("双模式", "快速/专业真实浏览器路径", "能选方案、能进入专业、同 projectId"),
    ("媒体", "导入/转写/配音/音乐/对齐", "真实源版本、词级时间戳、失败不写坏时间线"),
    ("预览导出", "合成画面、ffprobe、导出回读", "无黑帧/冻结帧/长静音、字幕安全区、音画同步"),
    ("管理边界", "SSO、积分、管理员供应商", "普通用户不配 API，管理员配置统一生效"),
    ("发布包", "版本哈希、回滚链接、AGPL 清单", "包不含媒体、.env、密钥；候选与版本链接一致"),
]
mt = doc.add_table(rows=1, cols=3)
mt.style = "Table Grid"
mt.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, h in enumerate(["范围", "真实检查", "通过条件"]):
    set_cell_text(mt.rows[0].cells[i], h, bold=True, color="FFFFFF", size=8)
    shade(mt.rows[0].cells[i], "374151")
for row in matrix:
    cells = mt.add_row().cells
    for i, value in enumerate(row):
        set_cell_text(cells[i], value, size=8)

doc.add_heading("Codex 总入口提示词", level=1)
add_body(doc, "你是念念智剪本阶段唯一主控。先读 AGENTS.md、DESIGN.md、ITERATION_PLAN_ZH.md、docs/FAST_PRO_MODE_PLAN_ZH.md、本 V2 文档和对应 Skill 完整章节。先报告：当前证据 -> 技能中文名 -> 具体章节 -> 验证或交付；只选一个主 Skill，最多一个支持 Skill。先输出用户能看懂的目标状态、真实输入、唯一交付物和验收证据，再改代码。保持念念智剪与念念 AI/智选宿主边界、同一 projectId/真实工程、SSO/积分/供应商/品牌/AGPL 边界。未验证事实先做只读检查；不把计划、HTTP 200、任务 ID或构建成功当成用户结果。阶段完成后自动继续到下一缺失交付；未通过时输出最小缺口，不发布线上。")

doc.add_heading("当前结论与下一动作", level=1)
add_body(doc, "本 V2 是前期权威准备，不代表 S1–S7 已实现；5 张目标图只负责方向。当前最小高价值动作是：本地冻结 S1 双模式页面合同和 S2 方案注册表的可视化验收标准，再实现一条真实“无声素材解说”快速工程，并从快速模式进入专业模式。")
add_body(doc, "未验证：本轮未调用付费 Provider 生成真实口播、ASR、BGM 或导出；未发布线上；Image2 图像上游返回 502，因此目标图采用当前编辑器风格的浏览器原型截图，并已在文档中标明。")

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
